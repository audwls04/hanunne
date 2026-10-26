"""
표준 형식(혜택 1건)
  id  번호          c   분야(복지/보건/교통/행정/기타)   n  정책명
  d   지원내용      m   남성 대상 여부(True/False)       f  여성 대상 여부
  a0  최소나이      a1  최대나이 (없으면 None)
  sd  시·도 (빈 문자열 = 전국)   sg  시·군·구 목록 (빈 목록 = 시·도 전체)
  q   자격조건      t   신청/이용 기간    todo 사용자가 해야 할 일
  how 신청 방법     at  신청 장소         org  담당 기관   tel 문의전화   u 세부 URL
"""
import csv
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from services import schema  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw" / "hanunne_dataset.xlsx"
OUT = ROOT / "data"

# 시·도 이름 통일
SIDO_NAME = {"서울시": "서울특별시", "서울특별시": "서울특별시", "경기도": "경기도", "인천광역시": "인천광역시"}
# 오타·잘못 들어간 지역 바로잡기
SGG_FIX = {"관천시": "과천시"}
SGG_SIDO_FIX = {"청주시": "충청북도"}  # 경기 시트에 들어간 충북 혜택


def text(v):
    """빈 칸은 '' 로, 나머지는 앞뒤 공백 제거한 문자열로."""
    return "" if v is None else str(v).strip()


def age(v):
    """65, '만 65세', '만65세 이상 (일부 만60세 이상)' → 65.  못 찾으면 None"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    m = re.search(r"(\d+)", str(v))
    return int(m.group(1)) if m else None


def age_range(v):
    """'만 54~74세' 처럼 범위로 적힌 경우 (54, 74)"""
    m = re.search(r"(\d+)\s*~\s*(\d+)", text(v))
    return (int(m.group(1)), int(m.group(2))) if m else (age(v), None)


def yes(v):
    """'해당', '남녀' 등이면 True"""
    return text(v) in ("해당", "남녀", "O", "o", "Y", "y")


# 시트별 열 위치 (0부터 셈)
# 서울,부천 / 경기: 번호 분야 정책명 지원내용 성별(남) 성별(여) 최소나이 최대나이 시,도 시,군,구 자격조건 기간 할일 방법 장소 기관 전화 URL
STD = dict(id=0, c=1, n=2, d=3, male=4, female=5, a0=6, a1=7, sd=8, sg=9, q=10, t=11, todo=12, how=13, at=14, org=15, tel=16, u=17)
# 인천: 순서 분야 정책명 지원내용 성별(남) 성별(여) 나이 시,도 시,군,구 자격요건 기간 할일 방법 장소 선정방법 기관 전화 URL
INCHEON = dict(id=0, c=1, n=2, d=3, male=4, female=5, age_text=6, sd=7, sg=8, q=9, t=10, todo=11, how=12, at=13, org=15, tel=16, u=17)
# 수원: 순서 분야 정책명 지원내용 성별 최소나이 최대나이 거주지 자격조건 기간 할일 방법 장소 기관 전화 URL
SUWON = dict(id=0, c=1, n=2, d=3, gender=4, a0=5, a1=6, place=7, q=8, t=9, todo=10, how=11, at=12, org=13, tel=14, u=15)

SHEETS = [("서울,부천", STD), ("인천광역시", INCHEON), ("수원", SUWON), ("경기", STD)]


def convert_row(r, col):
    g = lambda k: r[col[k]] if k in col and col[k] < len(r) else None
    if not g("n") or not g("id"):
        return None  # 정책명 없는 빈 행 / 메모 행은 건너뜀

    # 성별
    if "gender" in col:  # 수원: '남녀' 한 칸
        male = female = True
    else:
        male, female = yes(g("male")), yes(g("female"))
        if not male and not female:  # 경기 시트처럼 비어 있으면 모두 대상으로 봄
            male = female = True

    # 나이
    if "age_text" in col:
        a0, a1 = age_range(g("age_text"))
    else:
        a0, a1 = age(g("a0")), age(g("a1"))

    # 지역
    if "place" in col:  # 수원: '전국' 또는 '수원'
        sd, sgg = ("", "") if text(g("place")) == "전국" else ("경기도", "수원시")
    else:
        sd, sgg = SIDO_NAME.get(text(g("sd")), text(g("sd"))), text(g("sg"))
    sg = [] if sgg in ("", "전체", "전국") else [SGG_FIX.get(s.strip(), s.strip()) for s in sgg.split(",")]
    for s in sg:
        if s in SGG_SIDO_FIX:
            sd = SGG_SIDO_FIX[s]

    return dict(
        id=int(g("id")), c=text(g("c")) or "기타", n=text(g("n")), d=text(g("d")),
        m=male, f=female, a0=a0, a1=a1, sd=sd, sg=sg,
        q=text(g("q")), t=text(g("t")), todo=text(g("todo")), how=text(g("how")),
        at=text(g("at")), org=text(g("org")), tel=text(g("tel")), u=text(g("u")),
    )


def main():
    wb = openpyxl.load_workbook(RAW, read_only=True, data_only=True)

    # 1) 혜택
    policies = []
    for name, col in SHEETS:
        for r in list(wb[name].iter_rows(values_only=True))[1:]:
            p = convert_row(r, col)
            if p:
                policies.append(p)

    # 2) 단어장 (같은 단어가 여러 혜택에 반복되므로 하나로 합치고, 나온 혜택 번호를 모음)
    by_id = {p["id"]: p for p in policies}
    words = {}
    for no, hard, easy in list(wb["단어장"].iter_rows(values_only=True, max_col=3))[1:]:
        hard, easy = text(hard), text(easy)
        if not hard or not easy:
            continue
        w = words.setdefault(hard, {"w": hard, "e": easy, "p": []})
        if isinstance(no, (int, float)) and int(no) not in w["p"]:
            w["p"].append(int(no))
    for w in words.values():  # 단어 분야 = 그 단어가 나온 혜택들의 가장 흔한 분야
        cats = Counter(by_id[i]["c"] for i in w["p"] if i in by_id)
        w["c"] = cats.most_common(1)[0][0] if cats else "기타"
    words = list(words.values())

    # 3) 저장
    (OUT / "policies.json").write_text(json.dumps(policies, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / "words.json").write_text(json.dumps(words, ensure_ascii=False, indent=1), encoding="utf-8")
    with open(OUT / "policies.csv", "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.DictWriter(f, fieldnames=[col for _, col in schema.POLICY_COLUMNS])
        wr.writeheader()
        for p in policies:
            wr.writerow(schema.policy_to_row(p))
    with open(OUT / "words.csv", "w", newline="", encoding="utf-8-sig") as f:
        wr = csv.DictWriter(f, fieldnames=[col for _, col in schema.WORD_COLUMNS])
        wr.writeheader()
        for w in words:
            wr.writerow(schema.word_to_row(w))

    # 4) 데이터 점검 보고 — 비어 있는 칸을 알려줌
    print(f"혜택 {len(policies)}건, 단어 {len(words)}개 저장 완료 → {OUT}")
    for label, key in [("지원내용", "d"), ("최소나이", "a0"), ("문의전화", "tel")]:
        miss = [p["id"] for p in policies if not p[key] and p[key] != 0]
        print(f"  - {label} 비어 있음: {len(miss)}건")


if __name__ == "__main__":
    main()